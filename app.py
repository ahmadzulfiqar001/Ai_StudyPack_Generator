import os
import tempfile
import re
from xml.sax.saxutils import escape

import streamlit as st
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem

from ai_workflow import StudyPackWorkflow
from utils import read_uploaded_file, validate_material

st.set_page_config(
page_title="AI Study Pack Generator",
page_icon="📚",
layout="wide",
)

DEFAULT_MODEL = os.getenv(
"GROQ_MODEL",
"openai/gpt-oss-20b",
)

def get_api_key():
try:
key = st.secrets.get("GROQ_API_KEY", "")
if key:
return str(key).strip()
except Exception:
pass

```
return os.getenv("GROQ_API_KEY", "").strip()
```

def create_pdf_file(content, title):
temp_file = tempfile.NamedTemporaryFile(
delete=False,
suffix=".pdf",
)
temp_path = temp_file.name
temp_file.close()

```
styles = getSampleStyleSheet()

title_style = ParagraphStyle(
    "StudyTitle",
    parent=styles["Title"],
    fontSize=20,
    leading=24,
    spaceAfter=16,
)

heading_style = ParagraphStyle(
    "StudyHeading",
    parent=styles["Heading2"],
    fontSize=14,
    leading=18,
    spaceBefore=10,
    spaceAfter=6,
)

body_style = ParagraphStyle(
    "StudyBody",
    parent=styles["BodyText"],
    fontSize=10,
    leading=15,
    spaceAfter=6,
)

bullet_style = ParagraphStyle(
    "StudyBullet",
    parent=body_style,
    leftIndent=12,
    firstLineIndent=0,
)

story = [
    Paragraph(
        escape(title),
        title_style,
    )
]

lines = content.splitlines()

for line in lines:
    stripped = line.strip()

    if not stripped:
        story.append(Spacer(1, 4))
        continue

    if stripped.startswith("### "):
        text = stripped[4:].strip()
        story.append(
            Paragraph(
                escape(text),
                heading_style,
            )
        )
        continue

    if stripped.startswith("## "):
        text = stripped[3:].strip()
        story.append(
            Paragraph(
                escape(text),
                heading_style,
            )
        )
        continue

    if stripped.startswith("# "):
        text = stripped[2:].strip()
        story.append(
            Paragraph(
                escape(text),
                heading_style,
            )
        )
        continue

    if stripped.startswith("- ") or stripped.startswith("* "):
        text = stripped[2:].strip()
        text = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", text)
        text = re.sub(r"__(.*?)__", r"<b>\1</b>", text)
        text = escape(text, quote=False)
        text = text.replace("&lt;b&gt;", "<b>")
        text = text.replace("&lt;/b&gt;", "</b>")

        story.append(
            Paragraph(
                "• " + text,
                bullet_style,
            )
        )
        continue

    text = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", stripped)
    text = re.sub(r"__(.*?)__", r"<b>\1</b>", text)
    text = escape(text, quote=False)
    text = text.replace("&lt;b&gt;", "<b>")
    text = text.replace("&lt;/b&gt;", "</b>")

    story.append(
        Paragraph(
            text,
            body_style,
        )
    )

document = SimpleDocTemplate(
    temp_path,
    pagesize=A4,
    rightMargin=18 * mm,
    leftMargin=18 * mm,
    topMargin=18 * mm,
    bottomMargin=18 * mm,
)

document.build(story)

with open(temp_path, "rb") as file:
    pdf_bytes = file.read()

try:
    os.remove(temp_path)
except OSError:
    pass

return pdf_bytes
```

def create_word_file(content, title):
temp_file = tempfile.NamedTemporaryFile(
delete=False,
suffix=".docx",
)
temp_path = temp_file.name
temp_file.close()

```
document = Document()

document.add_heading(
    title,
    level=0,
)

for line in content.splitlines():
    stripped = line.strip()

    if not stripped:
        document.add_paragraph()
        continue

    if stripped.startswith("### "):
        document.add_heading(
            stripped[4:].strip(),
            level=3,
        )
        continue

    if stripped.startswith("## "):
        document.add_heading(
            stripped[3:].strip(),
            level=2,
        )
        continue

    if stripped.startswith("# "):
        document.add_heading(
            stripped[2:].strip(),
            level=1,
        )
        continue

    if stripped.startswith("- ") or stripped.startswith("* "):
        paragraph = document.add_paragraph(
            style="List Bullet",
        )
        text = stripped[2:].strip()

        parts = re.split(
            r"(\*\*.*?\*\*|__.*?__)",
            text,
        )

        for part in parts:
            if not part:
                continue

            if (
                part.startswith("**")
                and part.endswith("**")
            ):
                run = paragraph.add_run(
                    part[2:-2]
                )
                run.bold = True
            elif (
                part.startswith("__")
                and part.endswith("__")
            ):
                run = paragraph.add_run(
                    part[2:-2]
                )
                run.bold = True
            else:
                paragraph.add_run(part)

        continue

    paragraph = document.add_paragraph()

    parts = re.split(
        r"(\*\*.*?\*\*|__.*?__)",
        stripped,
    )

    for part in parts:
        if not part:
            continue

        if (
            part.startswith("**")
            and part.endswith("**")
        ):
            run = paragraph.add_run(
                part[2:-2]
            )
            run.bold = True
        elif (
            part.startswith("__")
            and part.endswith("__")
        ):
            run = paragraph.add_run(
                part[2:-2]
            )
            run.bold = True
        else:
            paragraph.add_run(part)

for paragraph in document.paragraphs:
    for run in paragraph.runs:
        run.font.size = Pt(10)

document.save(temp_path)

with open(temp_path, "rb") as file:
    word_bytes = file.read()

try:
    os.remove(temp_path)
except OSError:
    pass

return word_bytes
```

st.title("📚 AI Study Pack Generator")

st.caption(
"A multi-stage AI workflow for planning, generating, "
"assessing, reviewing, and refining personalized study material."
)

with st.sidebar:
st.header("⚙️ Study Settings")

```
subject = st.text_input(
    "Subject",
    placeholder="e.g. Machine Learning",
)

level = st.selectbox(
    "Student Level",
    [
        "School",
        "College",
        "University",
        "Professional",
    ],
    index=2,
)

pack_size = st.selectbox(
    "Pack Size",
    [
        "Quick",
        "Standard",
        "Detailed",
    ],
    index=1,
)

model = st.text_input(
    "AI Model",
    value=DEFAULT_MODEL,
    help="Groq model used for generating the study pack.",
)

st.divider()

api_key = get_api_key()

if api_key:
    st.success("Groq API key detected")
    st.caption(f"Model: {model}")
else:
    st.error("Groq API key not detected")
    st.info("Add GROQ_API_KEY in Streamlit Secrets.")

st.divider()

st.markdown(
    """
    **Workflow**

    1. Planning
    2. Content Generation
    3. Assessment
    4. Review
    5. Refinement
    """
)
```

st.subheader("📖 Study Material")

saved_material = st.session_state.get(
"material",
"",
)

material = st.text_area(
"Paste your notes or lecture material",
value=saved_material,
height=300,
placeholder="Paste the material you want to turn into a personalized study pack...",
)

uploaded = st.file_uploader(
"Or upload a file",
type=["pdf", "txt", "md", "csv"],
)

if uploaded:
temp_path = None

```
try:
    suffix = os.path.splitext(uploaded.name)[1]

    with tempfile.NamedTemporaryFile(
        delete=False,
        suffix=suffix,
    ) as temp_file:
        temp_file.write(uploaded.getbuffer())
        temp_path = temp_file.name

    uploaded_text = read_uploaded_file(
        temp_path
    )

    if st.button(
        "Use Uploaded Material",
        use_container_width=True,
    ):
        st.session_state["material"] = uploaded_text
        st.rerun()

except Exception as exc:
    st.error(
        f"Could not read the uploaded file: {exc}"
    )

finally:
    if temp_path and os.path.exists(temp_path):
        try:
            os.remove(temp_path)
        except OSError:
            pass
```

st.divider()

generate = st.button(
"Generate Personalized Study Pack",
type="primary",
use_container_width=True,
)

if generate:
try:
validate_material(material)
except ValueError as exc:
st.error(str(exc))
st.stop()

```
api_key = get_api_key()

if not api_key:
    st.error("Groq API key is missing.")
    st.info(
        "Add GROQ_API_KEY to Streamlit Secrets and restart the app."
    )
    st.stop()

try:
    workflow = StudyPackWorkflow(
        api_key=api_key,
        model=model,
    )
except Exception as exc:
    st.error(
        f"Could not initialize AI workflow: {exc}"
    )
    st.stop()

status = st.empty()
progress_bar = st.progress(0)

def update_progress(index, name):
    progress = min(
        max(float(index) / 5.0, 0.0),
        1.0,
    )

    progress_bar.progress(progress)
    status.info(
        f"Stage {index}/5 - {name}..."
    )

try:
    with st.spinner(
        "Running AI workflow with Groq..."
    ):
        result = workflow.run(
            subject=subject,
            level=level,
            material=material,
            pack_size=pack_size,
            progress=update_progress,
        )

except Exception as exc:
    st.error(
        f"Study pack generation failed: {exc}"
    )
    st.stop()

progress_bar.progress(1.0)

status.success(
    "Study pack generation complete!"
)

errors = result.get(
    "errors",
    [],
)

if errors:
    with st.expander("Workflow warnings"):
        for item in errors:
            if isinstance(item, dict):
                stage = item.get(
                    "stage",
                    "Unknown stage",
                )
                error = item.get(
                    "error",
                    "Unknown error",
                )
                st.warning(
                    f"{stage}: {error}"
                )
            else:
                st.warning(
                    str(item)
                )

tabs = st.tabs(
    [
        "Planning",
        "Content",
        "Assessment",
        "Review",
        "Final Pack",
    ]
)

with tabs[0]:
    planning = result.get(
        "planning"
    )

    if planning:
        st.json(planning)
    else:
        st.info(
            "Planning stage was not available."
        )

with tabs[1]:
    content = result.get(
        "content"
    )

    if content:
        st.json(content)
    else:
        st.info(
            "Content generation stage was not available."
        )

with tabs[2]:
    assessment = result.get(
        "assessment"
    )

    if assessment:
        st.json(assessment)
    else:
        st.info(
            "Assessment stage was not available."
        )

with tabs[3]:
    review = result.get(
        "review"
    )

    if review:
        if isinstance(review, dict):
            score = review.get(
                "quality_score"
            )

            if isinstance(
                score,
                (int, float),
            ):
                st.metric(
                    "Quality Score",
                    score,
                )

            st.json(review)
        else:
            st.write(review)
    else:
        st.info(
            "Review stage was not available."
        )

with tabs[4]:
    final_pack = result.get(
        "final_pack",
        "",
    )

    if final_pack:
        st.markdown(final_pack)

        safe_subject = re.sub(
            r"[^a-zA-Z0-9_-]+",
            "_",
            subject.strip(),
        ).strip("_")

        if not safe_subject:
            safe_subject = "study_pack"

        markdown_bytes = final_pack.encode(
            "utf-8"
        )

        pdf_bytes = create_pdf_file(
            final_pack,
            subject or "AI Study Pack",
        )

        word_bytes = create_word_file(
            final_pack,
            subject or "AI Study Pack",
        )

        col1, col2, col3 = st.columns(3)

        with col1:
            st.download_button(
                label="Download Markdown",
                data=markdown_bytes,
                file_name=f"{safe_subject}_study_pack.md",
                mime="text/markdown",
                use_container_width=True,
            )

        with col2:
            st.download_button(
                label="Download PDF",
                data=pdf_bytes,
                file_name=f"{safe_subject}_study_pack.pdf",
                mime="application/pdf",
                use_container_width=True,
            )

        with col3:
            st.download_button(
                label="Download Word",
                data=word_bytes,
                file_name=f"{safe_subject}_study_pack.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
            )

    else:
        st.warning(
            "Final study pack was not generated."
        )
```

else:
st.info(
"Enter study material, choose your settings, and click Generate Personalized Study Pack to start the five-stage AI workflow."
)

st.divider()

st.caption(
"AI Study Pack Generator | Streamlit | Groq | Multi-stage AI Workflow"
)
